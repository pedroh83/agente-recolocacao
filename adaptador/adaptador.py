# -*- coding: utf-8 -*-
"""
ADAPTADOR DE CURRÍCULO POR VAGA
--------------------------------
Lê uma vaga (arquivo .txt na pasta "vagas" ou texto copiado), compara com os
currículos base A e B, escolhe o mais aderente e pede à API Gemini (gratuita)
uma versão ajustada à vaga, sem inventar experiências.

Gera, na pasta "saida":
  - Curriculo_<Empresa>_<Cargo>.docx   -> currículo adaptado, pronto para revisar
  - Relatorio_<Empresa>_<Cargo>.txt    -> compatibilidade, palavras-chave, lacunas, alertas
e registra tudo em historico.csv (abre no Excel).

Uso:
  python adaptador.py              -> processa todas as vagas .txt da pasta "vagas"
  python adaptador.py --colar      -> usa o texto que está na área de transferência (Ctrl+C)
  python adaptador.py --teste      -> confere configuração, currículos e conexão com a API
"""

import csv
import datetime as dt
import json
import os
import re
import shutil
import sys
import time
import unicodedata
from pathlib import Path

# Acentos corretos no Prompt de Comando do Windows
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

PASTA = Path(__file__).resolve().parent
ARQ_CONFIG = PASTA / "config_adaptador.txt"
PASTA_BASE = PASTA / "curriculos_base"
PASTA_VAGAS = PASTA / "vagas"
PASTA_PROCESSADAS = PASTA_VAGAS / "processadas"
PASTA_SAIDA = PASTA / "saida"
ARQ_HISTORICO = PASTA / "historico.csv"

# Modelos tentados em ordem caso o configurado não exista mais
MODELOS_RESERVA = ["gemini-3.8-flash", "gemini-3.7-flash", "gemini-3.5-flash", "gemini-flash-latest",
                   "gemini-3.5-flash-lite", "gemini-3.1-flash-lite", "gemini-2.5-flash"]


# ---------------------------------------------------------------------------
# 1. CONFIGURAÇÃO
# ---------------------------------------------------------------------------
def ler_config():
    cfg = {
        "GEMINI_API_KEY": "",
        "MODELO": "gemini-3.8-flash",
        "NOME": "",
        "CONTATO": "",
        "IDIOMA": "português do Brasil",
        "PAUSA_ENTRE_VAGAS": "6",
    }
    if ARQ_CONFIG.exists():
        for linha in ARQ_CONFIG.read_text(encoding="utf-8-sig").splitlines():
            linha = linha.strip()
            if not linha or linha.startswith("#") or "=" not in linha:
                continue
            chave, valor = linha.split("=", 1)
            cfg[chave.strip().upper()] = valor.strip()
    # Variável de ambiente tem prioridade, se existir
    if os.environ.get("GEMINI_API_KEY"):
        cfg["GEMINI_API_KEY"] = os.environ["GEMINI_API_KEY"]
    return cfg


# ---------------------------------------------------------------------------
# 2. LEITURA DOS CURRÍCULOS BASE
# ---------------------------------------------------------------------------
def ler_docx(caminho):
    from docx import Document

    doc = Document(str(caminho))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabela in doc.tables:
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(dict.fromkeys(celulas)))
    return "\n".join(partes)


def ler_arquivo_texto(caminho):
    caminho = Path(caminho)
    if caminho.suffix.lower() == ".docx":
        return ler_docx(caminho)
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return caminho.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return caminho.read_text(encoding="utf-8", errors="ignore")


def encontrar_curriculo(letra):
    for ext in (".docx", ".txt", ".md"):
        for nome in (f"curriculo_{letra}{ext}", f"Curriculo_{letra}{ext}", f"curriculo_{letra.lower()}{ext}"):
            p = PASTA_BASE / nome
            if p.exists():
                return p
    return None


def carregar_curriculos():
    curriculos = {}
    for letra in ("A", "B"):
        p = encontrar_curriculo(letra)
        if p:
            texto = ler_arquivo_texto(p).strip()
            if len(texto) > 200:
                curriculos[letra] = texto
            else:
                print(f"  ! {p.name} parece vazio ou muito curto — verifique.")
    return curriculos


# ---------------------------------------------------------------------------
# 3. PONTUAÇÃO LOCAL (estilo ATS, sem IA)
# ---------------------------------------------------------------------------
STOPWORDS = set("""
a o as os um uma uns umas de do da dos das em no na nos nas por para pra com sem sob sobre
e ou que se ao aos à às é ser são será sera estar está esta este esse essa isso isto
como mais menos muito muita muitos muitas seu sua seus suas nosso nossa nossos nossas
você voce vocês voces ele ela eles elas nós nos eu me te lhe já ja também tambem
entre até ate após apos desde quando onde qual quais quem cada todo toda todos todas
outro outra outros outras ter tem têm possuir possui possuem vaga vagas empresa empresas
área area areas áreas trabalho trabalhar atuar atuação atuacao atividades responsabilidades
requisitos requisito desejável desejavel diferencial diferenciais experiência experiencia
conhecimento conhecimentos conhecer buscamos procuramos oferecemos benefícios beneficios
será sera ser sendo bem boa bom grande grandes novo nova novos novas nível nivel anos ano
the and or of to in for with on at by an be is are as from our your you we will
""".split())


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", texto.lower())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return texto


def termos_relevantes(texto, limite=40):
    palavras = re.findall(r"[a-z0-9][a-z0-9+#./-]{2,}", normalizar(texto))
    palavras = [w.strip(".-/") for w in palavras]
    palavras = [w for w in palavras if w not in STOPWORDS and len(w) > 2 and not w.isdigit()]
    contagem = {}
    for w in palavras:
        contagem[w] = contagem.get(w, 0) + 1
    # Bigramas (ex.: "gestao qualidade", "levantamento requisitos")
    for w1, w2 in zip(palavras, palavras[1:]):
        bg = f"{w1} {w2}"
        contagem[bg] = contagem.get(bg, 0) + 1
    ordenados = sorted(contagem.items(), key=lambda kv: (-kv[1], kv[0]))
    return [t for t, n in ordenados if n >= 1][:limite]


def pontuar(vaga, curriculo):
    termos = termos_relevantes(vaga)
    if not termos:
        return 0, [], []
    cv = normalizar(curriculo)
    achados = [t for t in termos if t in cv]
    faltando = [t for t in termos if t not in cv]
    return round(100 * len(achados) / len(termos)), achados, faltando


# ---------------------------------------------------------------------------
# 4. CHAMADA À API GEMINI
# ---------------------------------------------------------------------------
PROMPT = """Você é um especialista em recrutamento e em sistemas ATS (filtros automáticos de currículo).

TAREFA: adaptar um currículo para UMA vaga específica.

REGRAS OBRIGATÓRIAS (não quebre nenhuma — veracidade vale mais que compatibilidade):
1. Use SOMENTE fatos presentes nos currículos base. NUNCA invente empresas, cargos, datas, números,
   certificações, ferramentas ou resultados que não estejam escritos neles.
2. Você PODE: reescrever frases, reordenar experiências e tópicos, destacar o que é mais aderente,
   usar os termos exatos da vaga QUANDO correspondem a algo que a pessoa realmente fez,
   resumir ou omitir itens pouco relevantes para a vaga.
3. Mantenha todas as experiências profissionais com empresa, cargo e período iguais aos do currículo base
   (pode reduzir a 1 tópico as menos relevantes, mas não apague o emprego).
4. Requisitos da vaga que o candidato NÃO comprova vão para "lacunas", nunca para o currículo.
5. Escreva em {idioma}, tom profissional, verbos no passado para experiências anteriores,
   tópicos curtos (até 2 linhas), sem tabelas, ícones ou colunas (para passar no ATS).
6. Resumo profissional com 3 a 5 linhas, focado no cargo da vaga.
7. Cada tópico de uma experiência só pode descrever atividades que o currículo base atribui àquela MESMA
   experiência. NÃO transfira para um emprego ferramentas, métodos ou termos da vaga (ex.: indicadores,
   gargalos, KPIs, POP, fluxograma, Jira, Bizagi, compliance) que o base não cita naquele emprego.
   Se o termo existe em outro lugar do base (ex.: competências), ele pode ir só em "competencias".
8. NÃO acrescente níveis (ex.: "intermediário", "avançado", "fluente"), quantidades, nem verbos de liderança
   ("liderei", "coordenei", "gerenciei equipe") que não estejam escritos no base para aquela experiência.
   Idiomas e cursos devem ser copiados exatamente como estão no base.
   Também NÃO promova o papel: "acompanhei" ou "apoiei" no base não vira "conduzi", "ministrei" ou "liderei".
9. Prefira reaproveitar as frases do base com pequenos ajustes de ordem e vocabulário a reescrever do zero.
10. Toda experiência com data de término (ex.: "06/2025 - 07/2026") é passada: use SEMPRE verbos no passado
    ("Conduzi", "Participei"), mesmo que o currículo base esteja no presente.
11. Mantenha nas experiências as entregas de maior peso do base (implantações, migrações e recertificações de
    normas, projetos conduzidos), pois elas comprovam experiência com projetos, mesmo que a vaga não cite ISO.
    Ao resumir, preserve os detalhes que dão peso à entrega (versões das normas, "incluindo recertificação").

PASSO 1: compare a vaga com o CURRÍCULO A e o CURRÍCULO B e escolha a melhor base.
         (Pontuação local de palavras-chave, só como referência: A={nota_a}%, B={nota_b}%)
PASSO 2: produza o currículo adaptado a partir da base escolhida (pode aproveitar fatos da outra versão,
         pois ambas descrevem a mesma pessoa).

Responda APENAS com um JSON válido, exatamente neste formato:
{{
  "empresa": "nome da empresa da vaga ou 'Nao informado'",
  "cargo": "título da vaga",
  "versao_escolhida": "A ou B",
  "motivo_escolha": "1 a 2 frases",
  "compatibilidade_base": 0,
  "compatibilidade_adaptado": 0,
  "palavras_chave_vaga": ["até 15 termos mais importantes da vaga"],
  "palavras_chave_usadas": ["termos da vaga que entraram no currículo adaptado"],
  "lacunas": ["requisitos da vaga que o currículo não comprova"],
  "curriculo": {{
    "contato": "linha de contato copiada do currículo base (cidade | telefone | e-mail | LinkedIn)",
    "titulo": "título profissional alinhado à vaga, ex.: Analista de Requisitos | Processos | Qualidade",
    "resumo": "texto",
    "competencias": ["8 a 14 competências, priorizando as pedidas na vaga"],
    "experiencias": [
      {{"cargo": "", "empresa": "", "periodo": "", "topicos": [""]}}
    ],
    "formacao": ["Curso — Instituição (ano)"],
    "cursos_certificacoes": [""],
    "idiomas": [""]
  }},
  "dicas_candidatura": ["2 a 4 dicas práticas para esta candidatura (ex.: o que destacar na entrevista)"]
}}

=== VAGA ===
{vaga}

=== CURRÍCULO A (Qualidade & Processos) ===
{cv_a}

=== CURRÍCULO B (Requisitos & Tecnologia) ===
{cv_b}
"""


def extrair_json(texto):
    texto = texto.strip()
    texto = re.sub(r"^```(?:json)?\s*|\s*```$", "", texto)
    inicio, fim = texto.find("{"), texto.rfind("}")
    if inicio == -1 or fim == -1:
        raise ValueError("A IA não devolveu um JSON.")
    return json.loads(texto[inicio:fim + 1])


class Relogio:
    """Mostra na tela que o programa continua trabalhando enquanto espera a IA."""

    def __init__(self, a_cada=20):
        import threading
        self._parar = threading.Event()
        self._inicio = time.time()
        self._a_cada = a_cada
        self._t = threading.Thread(target=self._rodar, daemon=True)
        self._t.start()

    def _rodar(self):
        while not self._parar.wait(self._a_cada):
            print(f"    ... ainda aguardando a IA ({int(time.time() - self._inicio)}s)", flush=True)

    def parar(self):
        self._parar.set()


def chamar_gemini(prompt, cfg):
    from google import genai
    from google.genai import types

    if not cfg["GEMINI_API_KEY"] or "COLE_SUA_CHAVE" in cfg["GEMINI_API_KEY"]:
        raise RuntimeError("Chave da API não configurada. Abra config_adaptador.txt e preencha GEMINI_API_KEY.")

    limite = int(float(cfg.get("TEMPO_LIMITE_SEGUNDOS") or 120))
    cliente = genai.Client(api_key=cfg["GEMINI_API_KEY"],
                           http_options=types.HttpOptions(timeout=limite * 1000))
    # Começa pelo último modelo que respondeu bem (evita insistir num modelo sobrecarregado)
    arq_ok = PASTA / "ultimo_modelo_ok.txt"
    ultimo_ok = arq_ok.read_text(encoding="utf-8").strip() if arq_ok.exists() else ""
    ordem = [ultimo_ok, cfg["MODELO"]] + MODELOS_RESERVA
    modelos = [m for i, m in enumerate(ordem) if m and m not in ordem[:i]]
    ultimo_erro = None

    for modelo in modelos:
        for tentativa in range(1, 4):
            relogio = Relogio()
            try:
                resp = cliente.models.generate_content(
                    model=modelo,
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        temperature=0.3,
                        response_mime_type="application/json",
                        automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True),
                    ),
                )
                relogio.parar()
                resultado = extrair_json(resp.text or "")
                if modelo != cfg["MODELO"]:
                    print(f"  (resposta veio do modelo {modelo})")
                try:
                    arq_ok.write_text(modelo, encoding="utf-8")
                except Exception:
                    pass
                return resultado
            except Exception as e:  # noqa: BLE001
                relogio.parar()
                msg = str(e)
                ultimo_erro = e
                if "timed out" in msg.lower() or "timeout" in type(e).__name__.lower() or "ReadTimeout" in msg:
                    print(f"  ! {modelo} não respondeu em {limite}s. Tentando outro modelo...")
                    break
                if "404" in msg or "NOT_FOUND" in msg or "not found" in msg.lower():
                    print(f"  ! Modelo {modelo} indisponível, tentando o próximo...")
                    break  # vai para o próximo modelo
                if "429" in msg or "RESOURCE_EXHAUSTED" in msg:
                    if tentativa < 3:
                        espera = 30 * tentativa
                        print(f"  ! Limite gratuito atingido. Aguardando {espera}s (tentativa {tentativa}/3)...")
                        time.sleep(espera)
                        continue
                    print(f"  ! Limite do {modelo} esgotado. Tentando outro modelo...")
                    break
                if "API key" in msg or "API_KEY" in msg or "401" in msg or "403" in msg:
                    raise RuntimeError("Acesso negado pela API (chave inválida/errada, ou rede bloqueando). "
                                       "Confira a chave no config ou gere outra no Google AI Studio.") from e
                if isinstance(e, (ValueError, json.JSONDecodeError)) and tentativa < 3:
                    print("  ! Resposta fora do formato, pedindo de novo...")
                    continue
                if "500" in msg or "503" in msg or "UNAVAILABLE" in msg or "overloaded" in msg.lower():
                    if tentativa == 1:
                        print(f"  ! {modelo} ocupado. Tentando de novo em 8s...")
                        time.sleep(8)
                        continue
                    print(f"  ! {modelo} continua ocupado. Trocando de modelo...")
                    time.sleep(2)
                    break  # vai para o próximo modelo
                raise
    raise RuntimeError("O Google está sobrecarregado agora e nenhum modelo respondeu. "
                       f"Tente de novo daqui a alguns minutos com adaptar_vagas.bat. (Último erro: {str(ultimo_erro)[:120]})")


# ---------------------------------------------------------------------------
# 5. VERIFICAÇÃO DE VERACIDADE (a IA inventou algo?)
# ---------------------------------------------------------------------------
def verificar_veracidade(resultado, curriculos):
    base = normalizar("\n".join(curriculos.values()))
    base_compacta = re.sub(r"[^a-z0-9]", "", base)
    alertas = []
    cv = resultado.get("curriculo", {})
    for exp in cv.get("experiencias", []):
        empresa = exp.get("empresa", "")
        nucleo = re.sub(r"[^a-z0-9]", "", normalizar(empresa.split("(")[0].split("—")[0].split("-")[0]))
        # 10 letras: com 6, "Grupo Fantasia" passaria só porque o base tem um "Grupo ..."
        if nucleo and nucleo[:10] not in base_compacta:
            alertas.append(f"Empresa '{empresa}' não aparece nos currículos base.")
        texto_exp = " ".join(exp.get("topicos", []))
        for numero in set(re.findall(r"\d+(?:[.,]\d+)?\s*%?", texto_exp)):
            n = numero.strip()
            if n and n.replace(" ", "") not in base.replace(" ", ""):
                alertas.append(f"Número '{n}' em '{empresa}' não aparece nos currículos base — confira.")
    alertas += verificar_verbos(cv, curriculos)
    for item in cv.get("cursos_certificacoes", []) + cv.get("formacao", []) + cv.get("idiomas", []):
        chave = re.sub(r"[^a-z0-9]", "", normalizar(item))[:10]
        if chave and chave not in base_compacta:
            alertas.append(f"Formação/curso '{item}' não encontrado literalmente nos currículos base — confira.")
    return alertas


PRESENTE_PARA_PASSADO = {
    "conduzo": "conduzi", "participo": "participei", "executo": "executei", "valido": "validei",
    "realizo": "realizei", "documento": "documentei", "elaboro": "elaborei", "mapeio": "mapeei",
    "atuo": "atuei", "gerencio": "gerenciei", "coordeno": "coordenei", "lidero": "liderei",
    "acompanho": "acompanhei", "apoio": "apoiei", "desenvolvo": "desenvolvi", "analiso": "analisei",
    "implemento": "implementei", "planejo": "planejei", "ministro": "ministrei", "reviso": "revisei",
    "traduzo": "traduzi", "levanto": "levantei", "monitoro": "monitorei", "estruturo": "estruturei",
    "dão": "deram", "dá": "deu",
}


def corrigir_tempo_verbal(resultado):
    """Garante verbos no passado em experiências que já terminaram (ex.: '06/2025 - 07/2026')."""
    trocas = 0
    for exp in resultado.get("curriculo", {}).get("experiencias", []):
        periodo = normalizar(exp.get("periodo", ""))
        if not re.search(r"\d{2}/\d{4}\s*[-–a]+\s*\d{2}/\d{4}", periodo) or "atual" in periodo:
            continue
        novos = []
        for t in exp.get("topicos", []):
            def troca(m):
                w = m.group(0)
                passado = PRESENTE_PARA_PASSADO.get(w.lower())
                if not passado:
                    return w
                return passado.capitalize() if w[0].isupper() else passado
            # só troca no início do tópico ou logo após "e"/"que" (evita mexer em substantivos como "o documento")
            novo = re.sub(r"(?:(?<=^)|(?<=\be )|(?<=\bque ))[\wÀ-ÿ]+", troca, t)
            trocas += novo != t
            novos.append(novo)
        exp["topicos"] = novos
    return trocas


VERBOS_FORTES = ("liderei", "lidero", "coordenei", "coordeno", "ministrei", "ministro", "gerenciei",
                 "gerencio", "chefiei", "supervisionei", "implementei", "estruturei", "criei", "desenvolvi")


def trecho_da_empresa(texto_base, empresa):
    """Linhas do currículo base que pertencem a uma experiência (do cabeçalho até o próximo cabeçalho)."""
    nucleo = re.sub(r"[^a-z0-9]", "", normalizar(empresa.split("(")[0].split("—")[0]))[:6]
    if not nucleo:
        return ""
    linhas = texto_base.splitlines()
    for i, linha in enumerate(linhas):
        if nucleo in re.sub(r"[^a-z0-9]", "", normalizar(linha)) and re.search(r"\d{2}/\d{4}", linha):
            bloco = []
            for prox in linhas[i + 1:]:
                if re.search(r"\d{2}/\d{4}\s*-\s*\d{2}/\d{4}", prox) or prox.strip().isupper():
                    break
                bloco.append(prox)
            return normalizar(" ".join(bloco))
    return ""


def verificar_verbos(cv, curriculos):
    """Avisa quando um verbo de liderança/autoria aparece numa experiência sem estar no base daquela empresa."""
    alertas = []
    for exp in cv.get("experiencias", []):
        empresa = exp.get("empresa", "")
        blocos = " ".join(trecho_da_empresa(t, empresa) for t in curriculos.values())
        if not blocos.strip():
            continue
        for topico in exp.get("topicos", []):
            primeira = normalizar(topico).split()[:1]
            for verbo in VERBOS_FORTES:
                if (primeira and primeira[0] == verbo or f" {verbo} " in f" {normalizar(topico)} ") and verbo not in blocos:
                    alertas.append(f"'{verbo}' em '{empresa}' não aparece no base dessa experiência — confira: "
                                   f"\"{topico[:80]}...\"")
                    break
    return alertas


# ---------------------------------------------------------------------------
# 6. GERAÇÃO DO .DOCX (layout simples, amigável a ATS)
# ---------------------------------------------------------------------------
def gerar_docx(resultado, cfg, destino):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    cv = resultado["curriculo"]
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.0)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    estilo.font.size = Pt(10.5)
    estilo.paragraph_format.space_after = Pt(2)

    def paragrafo(texto="", negrito=False, tamanho=None, alinhamento=None, cor=None, depois=None):
        p = doc.add_paragraph()
        if texto:
            r = p.add_run(texto)
            r.bold = negrito
            if tamanho:
                r.font.size = Pt(tamanho)
            if cor:
                r.font.color.rgb = RGBColor.from_string(cor)
        if alinhamento is not None:
            p.alignment = alinhamento
        if depois is not None:
            p.paragraph_format.space_after = Pt(depois)
        return p

    def titulo_secao(texto):
        p = paragrafo(texto.upper(), negrito=True, tamanho=11.5, cor="1F3864", depois=3)
        p.paragraph_format.space_before = Pt(10)
        pPr = p._p.get_or_add_pPr()
        borda = OxmlElement("w:pBdr")
        inferior = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3864")):
            inferior.set(qn(k), v)
        borda.append(inferior)
        pPr.append(borda)

    def topico(texto):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(texto)
        p.paragraph_format.space_after = Pt(1)

    nome = cfg.get("NOME") or "Seu Nome"
    contato = cfg.get("CONTATO") or cv.get("contato", "")
    paragrafo(nome, negrito=True, tamanho=18, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, cor="1F3864", depois=0)
    if cv.get("titulo"):
        paragrafo(cv["titulo"], negrito=True, tamanho=11.5, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, depois=0)
    if contato:
        paragrafo(contato, tamanho=10, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, depois=4)

    if cv.get("resumo"):
        titulo_secao("Resumo Profissional")
        p = paragrafo(cv["resumo"])
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if cv.get("competencias"):
        titulo_secao("Competências")
        paragrafo(" • ".join(cv["competencias"]))

    if cv.get("experiencias"):
        titulo_secao("Experiência Profissional")
        for exp in cv["experiencias"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(exp.get("cargo", ""))
            r.bold = True
            if exp.get("empresa"):
                p.add_run(f" — {exp['empresa']}")
            if exp.get("periodo"):
                r2 = p.add_run(f"  |  {exp['periodo']}")
                r2.italic = True
            for t in exp.get("topicos", []):
                if t.strip():
                    topico(t)

    for chave, rotulo in (("formacao", "Formação Acadêmica"),
                          ("cursos_certificacoes", "Cursos e Certificações"),
                          ("idiomas", "Idiomas")):
        itens = [i for i in cv.get(chave, []) if str(i).strip()]
        if itens:
            titulo_secao(rotulo)
            for i in itens:
                topico(i)

    doc.core_properties.author = nome
    doc.core_properties.title = f"Currículo — {resultado.get('cargo', '')}"
    doc.save(str(destino))


# ---------------------------------------------------------------------------
# 7. RELATÓRIO E HISTÓRICO
# ---------------------------------------------------------------------------
def gerar_relatorio(resultado, notas_locais, alertas, link, arq_vaga, destino):
    L = []
    L.append("RELATÓRIO DE ADAPTAÇÃO DE CURRÍCULO")
    L.append("=" * 60)
    L.append(f"Data: {dt.datetime.now():%d/%m/%Y %H:%M}")
    L.append(f"Vaga: {resultado.get('cargo', '')}  |  Empresa: {resultado.get('empresa', '')}")
    if link:
        L.append(f"Link: {link}")
    L.append(f"Arquivo da vaga: {arq_vaga}")
    L.append("")
    L.append(f"Versão base escolhida: {resultado.get('versao_escolhida', '?')}")
    L.append(f"Motivo: {resultado.get('motivo_escolha', '')}")
    L.append(f"Pontuação local de palavras-chave: A = {notas_locais.get('A', '-')}%  |  B = {notas_locais.get('B', '-')}%")
    L.append(f"Compatibilidade estimada pela IA: base {resultado.get('compatibilidade_base', '?')}%  ->  "
             f"adaptado {resultado.get('compatibilidade_adaptado', '?')}%")
    L.append("")
    L.append("PALAVRAS-CHAVE DA VAGA")
    L.append("  " + ", ".join(resultado.get("palavras_chave_vaga", [])))
    L.append("PALAVRAS-CHAVE USADAS NO CURRÍCULO")
    L.append("  " + ", ".join(resultado.get("palavras_chave_usadas", [])))
    L.append("")
    L.append("LACUNAS (requisitos que seu currículo não comprova — NÃO foram inseridos)")
    for item in resultado.get("lacunas", []) or ["Nenhuma apontada."]:
        L.append(f"  - {item}")
    L.append("")
    L.append("ALERTAS DE VERACIDADE (confira antes de enviar)")
    for item in alertas or ["Nenhum: empresas, números e cursos batem com os currículos base."]:
        L.append(f"  - {item}")
    L.append("")
    L.append("DICAS PARA ESTA CANDIDATURA")
    for item in resultado.get("dicas_candidatura", []):
        L.append(f"  - {item}")
    L.append("")
    L.append("ANTES DE ENVIAR: abra o .docx, leia tudo, ajuste o que quiser e salve como PDF se a vaga pedir.")
    Path(destino).write_text("\n".join(L), encoding="utf-8-sig")


def registrar_historico(linha):
    novo = not ARQ_HISTORICO.exists()
    with open(ARQ_HISTORICO, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        if novo:
            w.writerow(["Data", "Empresa", "Cargo", "Versão", "Nota local A", "Nota local B",
                        "Compat. base IA", "Compat. adaptado IA", "Alertas", "Link", "Currículo gerado"])
        w.writerow(linha)


# ---------------------------------------------------------------------------
# 8. FLUXO PRINCIPAL
# ---------------------------------------------------------------------------
def limpar_nome(texto, limite=40):
    texto = normalizar(texto or "")
    texto = re.sub(r"[^a-z0-9]+", "_", texto).strip("_")
    return (texto[:limite].strip("_") or "vaga").title()


def extrair_link(texto):
    m = re.search(r"https?://\S+", texto)
    return m.group(0).rstrip(").,;") if m else ""


def processar_vaga(arq_vaga, curriculos, cfg, chamar=None):
    chamar = chamar or chamar_gemini
    texto_vaga = ler_arquivo_texto(arq_vaga).strip()
    if len(texto_vaga) < 150:
        print(f"  ! {arq_vaga.name}: texto muito curto (copie a descrição completa da vaga). Pulando.")
        return None
    link = extrair_link(texto_vaga)

    notas, faltando = {}, {}
    for letra, cv in curriculos.items():
        notas[letra], _, faltando[letra] = pontuar(texto_vaga, cv)
    print("  Pontuação local: " + "  ".join(f"{k}={v}%" for k, v in notas.items()))

    prompt = PROMPT.format(
        idioma=cfg["IDIOMA"],
        nota_a=notas.get("A", "-"), nota_b=notas.get("B", "-"),
        vaga=texto_vaga[:15000],
        cv_a=curriculos.get("A", "(não fornecido)"),
        cv_b=curriculos.get("B", "(não fornecido)"),
    )
    print("  Consultando a IA (Gemini)...")
    resultado = chamar(prompt, cfg)

    corrigir_tempo_verbal(resultado)
    alertas = verificar_veracidade(resultado, curriculos)

    base_nome = f"{limpar_nome(resultado.get('empresa'), 25)}_{limpar_nome(resultado.get('cargo'), 35)}"
    carimbo = dt.datetime.now().strftime("%Y%m%d_%H%M")
    PASTA_SAIDA.mkdir(exist_ok=True)
    arq_docx = PASTA_SAIDA / f"Curriculo_{base_nome}_{carimbo}.docx"
    arq_rel = PASTA_SAIDA / f"Relatorio_{base_nome}_{carimbo}.txt"

    gerar_docx(resultado, cfg, arq_docx)
    gerar_relatorio(resultado, notas, alertas, link, arq_vaga.name, arq_rel)
    registrar_historico([
        dt.datetime.now().strftime("%d/%m/%Y %H:%M"), resultado.get("empresa", ""), resultado.get("cargo", ""),
        resultado.get("versao_escolhida", ""), notas.get("A", ""), notas.get("B", ""),
        resultado.get("compatibilidade_base", ""), resultado.get("compatibilidade_adaptado", ""),
        len(alertas), link, arq_docx.name,
    ])

    print(f"  OK  Versão {resultado.get('versao_escolhida')} | compatibilidade "
          f"{resultado.get('compatibilidade_base')}% -> {resultado.get('compatibilidade_adaptado')}%")
    print(f"      {arq_docx.name}")
    if alertas:
        print(f"      ATENÇÃO: {len(alertas)} alerta(s) de veracidade — veja o relatório.")
    return arq_docx


def texto_da_area_de_transferencia():
    try:
        import tkinter
        raiz = tkinter.Tk()
        raiz.withdraw()
        texto = raiz.clipboard_get()
        raiz.destroy()
        return texto
    except Exception:
        return ""


def preparar_pastas():
    for p in (PASTA_BASE, PASTA_VAGAS, PASTA_PROCESSADAS, PASTA_SAIDA):
        p.mkdir(parents=True, exist_ok=True)


def modo_teste(cfg, curriculos):
    print("TESTE DE CONFIGURAÇÃO")
    print(f"  Pasta: {PASTA}")
    print(f"  Config encontrado: {'sim' if ARQ_CONFIG.exists() else 'NÃO'}")
    chave = cfg["GEMINI_API_KEY"]
    print(f"  Chave da API: {'preenchida (' + chave[:6] + '...)' if chave and 'COLE' not in chave else 'NÃO preenchida'}")
    print(f"  Modelo: {cfg['MODELO']}")
    for letra in ("A", "B"):
        p = encontrar_curriculo(letra)
        print(f"  Currículo {letra}: {p.name + ' (' + str(len(curriculos.get(letra, ''))) + ' caracteres)' if p else 'NÃO encontrado'}")
    try:
        import docx  # noqa: F401
        from google import genai  # noqa: F401
        print("  Bibliotecas: OK")
    except ImportError as e:
        print(f"  Bibliotecas: FALTANDO ({e}). Rode instalar.bat")
        return
    if chave and "COLE" not in chave:
        try:
            from google import genai
            from google.genai import types
            cliente = genai.Client(api_key=chave)
            r = cliente.models.generate_content(
                model=cfg["MODELO"], contents="Responda apenas: OK",
                config=types.GenerateContentConfig(
                    automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True)),
            )
            print(f"  Conexão com a API: OK (resposta: {r.text.strip()[:20]})")
        except Exception as e:  # noqa: BLE001
            print(f"  Conexão com a API: FALHOU -> {str(e)[:200]}")


# ---------------------------------------------------------------------------
# 9. ASSISTENTE DE INSTALAÇÃO (python adaptador.py --configurar)
# ---------------------------------------------------------------------------
def salvar_config_valor(chave, valor):
    linhas = ARQ_CONFIG.read_text(encoding="utf-8-sig").splitlines() if ARQ_CONFIG.exists() else []
    trocou = False
    for i, linha in enumerate(linhas):
        if linha.strip().upper().startswith(chave + "="):
            linhas[i] = f"{chave}={valor}"
            trocou = True
    if not trocou:
        linhas.append(f"{chave}={valor}")
    ARQ_CONFIG.write_text("\n".join(linhas) + "\n", encoding="utf-8")


def perguntar(texto, padrao=""):
    sufixo = f" [{padrao}]" if padrao else ""
    resposta = input(f"{texto}{sufixo}: ").strip().strip('"')
    return resposta or padrao


def abrir(caminho_ou_url):
    try:
        if str(caminho_ou_url).startswith("http"):
            import webbrowser
            webbrowser.open(str(caminho_ou_url))
        elif hasattr(os, "startfile"):
            os.startfile(str(caminho_ou_url))  # Windows: abre no Explorador
    except Exception:
        pass


def testar_chave(chave, modelo):
    """Devolve (ok, mensagem)."""
    from google import genai
    from google.genai import types
    cliente = genai.Client(api_key=chave)
    ultimo = ""
    for m in [modelo] + [x for x in MODELOS_RESERVA if x != modelo]:
        try:
            cliente.models.generate_content(
                model=m, contents="Responda apenas: OK",
                config=types.GenerateContentConfig(
                    automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True)),
            )
            if m != modelo:
                salvar_config_valor("MODELO", m)
            return True, m
        except Exception as e:  # noqa: BLE001
            ultimo = str(e)
            if "404" in ultimo or "NOT_FOUND" in ultimo:
                continue
            return False, ultimo[:200]
    return False, ultimo[:200]


def pasta_area_de_trabalho():
    """Pasta real da Área de Trabalho (funciona também quando ela está no OneDrive)."""
    try:
        import ctypes
        from ctypes import wintypes
        buf = ctypes.create_unicode_buffer(wintypes.MAX_PATH)
        # 0x10 = CSIDL_DESKTOPDIRECTORY
        if ctypes.windll.shell32.SHGetFolderPathW(None, 0x10, None, 0, buf) == 0 and Path(buf.value).exists():
            return Path(buf.value)
    except Exception:
        pass
    for p in (Path.home() / "OneDrive" / "Área de Trabalho", Path.home() / "OneDrive" / "Desktop",
              Path.home() / "Desktop"):
        if p.exists():
            return p
    return None


def criar_atalho():
    desktop = pasta_area_de_trabalho()
    if not desktop:
        print("  Não encontrei a Área de Trabalho. Use o colar_vaga.bat direto da pasta Adaptador.")
        return None
    atalho = desktop / "Adaptar Curriculo.bat"
    atalho.write_text(f'@call "{PASTA / "colar_vaga.bat"}"\r\n', encoding="cp1252", errors="ignore")
    print(f"  Atalho criado: {atalho}")
    return atalho


def assistente():
    linha = "-" * 60
    print(linha)
    print(" ASSISTENTE DE INSTALAÇÃO DO ADAPTADOR DE CURRÍCULO")
    print(linha)
    print(" São 4 etapas. Responda às perguntas e aperte Enter.\n")

    # Etapa 1 — pastas
    preparar_pastas()
    print("ETAPA 1 de 4 — Pastas criadas. OK\n")

    # Etapa 2 — chave da API
    cfg = ler_config()
    chave_atual = cfg["GEMINI_API_KEY"]
    tem_chave = chave_atual and "COLE" not in chave_atual
    print("ETAPA 2 de 4 — Chave gratuita do Google Gemini")
    if tem_chave and perguntar("  Já existe uma chave salva. Quer trocar? (s/N)", "n").lower() != "s":
        print("  Mantida a chave atual.\n")
    else:
        print("  Vou abrir a página do Google AI Studio no navegador.")
        print("  Lá: entre com sua conta Google > clique em 'Create API key' > clique no ícone de copiar.")
        input("  Aperte Enter para abrir a página...")
        abrir("https://aistudio.google.com/apikey")
        for tentativa in range(1, 4):
            print()
            chave = perguntar("  Cole a chave aqui (clique com o botão direito para colar) e aperte Enter")
            if len(chave) < 20:
                print("  Isso não parece uma chave. Copie de novo no AI Studio.")
                continue
            print("  Testando a chave...")
            ok, info = testar_chave(chave, cfg["MODELO"])
            if ok:
                salvar_config_valor("GEMINI_API_KEY", chave)
                print(f"  Chave funcionando (modelo {info}). Salva!\n")
                break
            print(f"  A chave não funcionou: {info}")
            if tentativa == 3:
                salvar_config_valor("GEMINI_API_KEY", chave)
                print("  Salvei mesmo assim. Confira a chave depois e rode instalar.bat de novo.\n")

    # Etapa 3 — dados do cabeçalho
    print("ETAPA 3 de 4 — Seus dados para o topo do currículo (Enter mantém o que está entre colchetes)")
    cfg = ler_config()
    # Usa o que já está salvo como sugestão, para não apagar nada ao rodar de novo
    partes = [x.strip() for x in cfg.get("CONTATO", "").split("|") if x.strip() and "SEU-PERFIL" not in x]
    atual_email = next((x for x in partes if "@" in x), "")
    atual_link = next((x for x in partes if "linkedin" in x.lower()), "")
    atual_tel = next((x for x in partes if sum(c.isdigit() for c in x) >= 8 and x not in (atual_email, atual_link)), "")
    if "00000-0000" in atual_tel:
        atual_tel = ""
    atual_cidade = next((x for x in partes if x not in (atual_email, atual_link, atual_tel)), "")
    nome = perguntar("  Nome completo", cfg.get("NOME") or "")
    cidade = perguntar("  Cidade - UF", atual_cidade)
    telefone = perguntar("  Telefone com DDD, ex.: (71) 99999-9999", atual_tel)
    email = perguntar("  E-mail", atual_email)
    linkedin = perguntar("  Link do LinkedIn, ex.: linkedin.com/in/seu-nome (Enter para pular)", atual_link)
    contato = " | ".join(x for x in (cidade, telefone, email, linkedin) if x)
    salvar_config_valor("NOME", nome)
    salvar_config_valor("CONTATO", contato)
    print(f"  Salvo: {nome} — {contato}\n")

    # Etapa 4 — currículos A e B
    print("ETAPA 4 de 4 — Seus currículos A e B")
    print("  Vou abrir a pasta 'curriculos_base'. Copie para ela seus dois currículos (.docx ou .txt),")
    print("  com o nome que tiverem. Depois volte para esta janela.")
    input("  Aperte Enter para abrir a pasta...")
    abrir(PASTA_BASE)
    usados = set()
    if not (encontrar_curriculo("A") and encontrar_curriculo("B")):
        input("\n  Copiou os arquivos? Aperte Enter para continuar...")
    for letra, descricao in (("A", "Qualidade & Processos"), ("B", "Requisitos & Tecnologia")):
        while not encontrar_curriculo(letra):
            arquivos = sorted(p for p in PASTA_BASE.iterdir()
                              if p.suffix.lower() in (".docx", ".txt", ".md") and not p.name.startswith("~$")
                              and not re.match(r"curriculo_[ab]\.", p.name.lower()))
            if not arquivos:
                print("  Nenhum .docx ou .txt na pasta ainda. Se o currículo for PDF, abra no Word e salve como .docx.")
                input("  Copie o arquivo e aperte Enter...")
                continue
            print(f"\n  Qual destes é o currículo {letra} ({descricao})?")
            for i, p in enumerate(arquivos, 1):
                marca = "   (já escolhido como A)" if p.name in usados else ""
                print(f"    {i}) {p.name}{marca}")
            escolha = perguntar("  Digite o número")
            if escolha.isdigit() and 1 <= int(escolha) <= len(arquivos) and arquivos[int(escolha) - 1].name in usados:
                print("  Esse já é o currículo A. Escolha o outro.")
            elif escolha.isdigit() and 1 <= int(escolha) <= len(arquivos):
                origem = arquivos[int(escolha) - 1]
                shutil.copy2(origem, PASTA_BASE / f"curriculo_{letra}{origem.suffix.lower()}")
                usados.add(origem.name)
                print(f"  OK: {origem.name} agora é o currículo {letra}.")
            else:
                print("  Número inválido, vamos de novo.")
    curriculos = carregar_curriculos()
    print(f"  Currículos prontos: {', '.join(sorted(curriculos))}\n")

    # Atalho na Área de Trabalho
    if perguntar("Criar o atalho 'Adaptar Curriculo' na Área de Trabalho? (S/n)", "s").lower() != "n":
        criar_atalho()
        print()

    print(linha)
    print(" TUDO PRONTO!")
    print(" Para usar: abra a vaga no LinkedIn, selecione a descrição, Ctrl+C")
    print(" e dê dois cliques em 'Adaptar Curriculo' (ou colar_vaga.bat).")
    print(linha)


def main():
    args = [a.lower() for a in sys.argv[1:]]
    if "--configurar" in args:
        assistente()
        return
    if "--atalho" in args:
        criar_atalho()
        return

    preparar_pastas()
    cfg = ler_config()
    curriculos = carregar_curriculos()

    if "--teste" in args:
        modo_teste(cfg, curriculos)
        return

    if not curriculos:
        print(f"Nenhum currículo base encontrado em {PASTA_BASE}")
        print("Coloque curriculo_A.docx e curriculo_B.docx (ou .txt) nessa pasta.")
        return

    if "--colar" in args:
        texto = texto_da_area_de_transferencia()
        if len(texto.strip()) < 150:
            print("A área de transferência está vazia ou com pouco texto.")
            print("Abra a vaga, selecione a descrição inteira, Ctrl+C e rode de novo.")
            return
        marcas = ("Pontuação local", "Consultando a IA", "Vaga copiada salva", "Traceback", "TUDO PRONTO",
                  "Pressione qualquer tecla", "ASSISTENTE DE INSTALAÇÃO")
        if any(m in texto for m in marcas):
            print("O texto copiado NÃO é uma vaga: parece ser texto desta janela preta.")
            print("Volte ao LinkedIn, selecione a descrição da vaga, aperte Ctrl+C e abra o atalho de novo.")
            return
        primeira = next((l.strip() for l in texto.splitlines() if l.strip()), "")[:90]
        print(f"Texto copiado começa com: \"{primeira}\"")
        arq = PASTA_VAGAS / f"vaga_{dt.datetime.now():%Y%m%d_%H%M%S}.txt"
        arq.write_text(texto, encoding="utf-8")
        print(f"Vaga copiada salva em {arq.name}")

    vagas = sorted(p for p in PASTA_VAGAS.glob("*.txt") if p.is_file())
    if not vagas:
        print(f"Nenhuma vaga para processar em {PASTA_VAGAS}")
        print("Salve a descrição da vaga como .txt nessa pasta, ou copie o texto e use colar_vaga.bat")
        return

    print(f"{len(vagas)} vaga(s) para adaptar.\n")
    pausa = float(cfg.get("PAUSA_ENTRE_VAGAS") or 6)
    gerados = 0
    for i, arq in enumerate(vagas, 1):
        print(f"[{i}/{len(vagas)}] {arq.name}")
        try:
            if processar_vaga(arq, curriculos, cfg):
                gerados += 1
                shutil.move(str(arq), str(PASTA_PROCESSADAS / arq.name))
        except Exception as e:  # noqa: BLE001
            print(f"  ERRO: {e}")
            print("  A vaga ficou na pasta 'vagas' para tentar de novo depois.")
        if i < len(vagas):
            time.sleep(pausa)
        print()

    print(f"Concluído: {gerados} currículo(s) gerado(s) em {PASTA_SAIDA}")


if __name__ == "__main__":
    main()
