# -*- coding: utf-8 -*-
"""
Controle unico de candidaturas (controle_candidaturas.xlsx).

Junta numa planilha so:
  - todas as vagas que o agente ja mostrou (historico vagas_vistas.json);
  - o "Enviado?" marcado nas planilhas do dia (vagas_AAAA-MM-DD.xlsx);
  - as adaptacoes feitas pelo Adaptador de Curriculo (Adaptador/historico.csv);
  - o que voce preencheu no proprio controle (Enviado?, Data do envio,
    Retorno e Observacoes), que nunca e apagado.

A aba "Indicadores" usa formulas do Excel e se atualiza sozinha quando voce
preenche o Retorno de uma candidatura.

A planilha e reconstruida a cada execucao a partir dessas fontes. Se ela
estiver aberta no Excel, nada se perde: a atualizacao fica para a proxima vez.
"""
from __future__ import annotations

import csv
import datetime as dt
import difflib
import glob
import hashlib
import os
import re
import shutil
from typing import Dict, List, Optional, Tuple

from configuracao import PalavrasChave
from extracao import id_da_vaga, link_limpo
from pontuacao import avaliar
from texto import normalizar

NOME_ARQUIVO = "controle_candidaturas.xlsx"
NOME_BACKUP = "controle_candidaturas_backup.xlsx"
LINHAS_FORMULA = 5000  # alcance das formulas e das listas de opcoes

# (titulo da coluna, largura, preenchida por voce?)
COLUNAS = [
    ("Data", 11, False),
    ("Nota", 6, False),
    ("Currículo", 10, False),
    ("Vaga", 44, False),
    ("Empresa", 28, False),
    ("Local", 22, False),
    ("Link", 11, False),
    ("Origem", 11, False),
    ("Adaptado?", 10, False),
    ("Compat. IA", 10, False),
    ("Alertas", 8, False),
    ("Enviado?", 10, True),
    ("Data do envio", 13, True),
    ("Retorno", 14, True),
    ("Observações", 36, True),
    ("ID", 16, False),
]
TITULOS = [c[0] for c in COLUNAS]
COL = {nome: i + 1 for i, nome in enumerate(TITULOS)}
CAMPOS_DO_USUARIO = [c[0] for c in COLUNAS if c[2]]

OPCOES_ENVIADO = ["Sim", "Não"]
OPCOES_RETORNO = ["Aguardando", "Sem retorno", "Entrevista", "Teste/Case", "Proposta", "Recusado"]
RETORNOS_POSITIVOS = ["Entrevista", "Teste/Case", "Proposta"]

RESPOSTAS_SIM = {"sim", "s", "x", "ok", "enviado", "enviada", "yes", "y"}
RESPOSTAS_NAO = {"nao", "n", "no"}


# ---------------------------------------------------------------------------
# Conversoes
# ---------------------------------------------------------------------------

def para_data(valor) -> Optional[dt.date]:
    """Aceita date, datetime, 'AAAA-MM-DD', 'DD/MM/AAAA' (com ou sem hora)."""
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(texto[:16] if "%H" in formato else texto[:10], formato).date()
        except ValueError:
            continue
    return None


def normalizar_enviado(valor) -> str:
    texto = normalizar(str(valor or "")).strip()
    if texto in RESPOSTAS_SIM:
        return "Sim"
    if texto in RESPOSTAS_NAO:
        return "Não"
    return str(valor).strip() if valor else ""


def _nucleo_empresa(nome: str) -> str:
    return re.sub(r"[^a-z0-9]", "", normalizar(nome or ""))


def mesma_empresa(a: str, b: str) -> bool:
    na, nb = _nucleo_empresa(a), _nucleo_empresa(b)
    if not na or not nb:
        return False
    return na[:6] in nb or nb[:6] in na


def semelhanca_titulo(a: str, b: str) -> float:
    na, nb = " ".join(normalizar(a).split()), " ".join(normalizar(b).split())
    if not na or not nb:
        return 0.0
    if na in nb or nb in na:
        return 1.0
    return difflib.SequenceMatcher(None, na, nb).ratio()


# ---------------------------------------------------------------------------
# Leitura das fontes
# ---------------------------------------------------------------------------

def ler_controle_existente(caminho: str) -> Tuple[Dict[str, dict], Dict[str, dict]]:
    """
    Devolve ({id: campos do usuario}, {id: linha completa}).

    A linha completa serve para nunca perder uma vaga que ja esta no controle,
    mesmo que ela saia das fontes (ex.: historico apagado) ou tenha sido
    incluida a mao (sem ID: recebe um ID "manual-...").
    """
    if not os.path.exists(caminho):
        return {}, {}
    from openpyxl import load_workbook

    wb = load_workbook(caminho)
    if "Candidaturas" not in wb.sheetnames:
        return {}, {}
    ws = wb["Candidaturas"]
    titulos = [str(c.value or "").strip() for c in ws[1]]
    pos = {t: i for i, t in enumerate(titulos) if t}

    usuario, existentes = {}, {}
    for linha in ws.iter_rows(min_row=2):
        def valor(nome):
            i = pos.get(nome)
            return linha[i].value if i is not None and i < len(linha) else None

        vaga, empresa = valor("Vaga"), valor("Empresa")
        ident = str(valor("ID") or "").strip()
        if not ident and not vaga:
            continue  # linha vazia
        if not ident:
            ident = "manual-" + hashlib.md5(normalizar(f"{vaga}|{empresa}").encode()).hexdigest()[:10]
        campos = {nome: valor(nome) for nome in CAMPOS_DO_USUARIO}

        link = ""
        i_link = pos.get("Link")
        if i_link is not None and i_link < len(linha):
            celula = linha[i_link]
            link = (celula.hyperlink.target if celula.hyperlink else None) or str(celula.value or "")
        compat = valor("Compat. IA")
        existentes[ident] = {
            "id": ident, "data": para_data(valor("Data")), "nota": valor("Nota"),
            "curriculo": valor("Currículo") or "", "titulo": str(vaga or ""),
            "empresa": str(empresa or ""), "local": str(valor("Local") or ""),
            "link": link if link.startswith("http") else "",
            "origem": "Manual" if ident.startswith("manual-") else (valor("Origem") or ""),
            "adaptado": valor("Adaptado?") or "",
            "compat": int(round(compat * 100)) if isinstance(compat, (int, float)) else None,
            "alertas": valor("Alertas"),
        }
        usuario[ident] = campos
    return usuario, existentes


def ler_enviados_das_planilhas_do_dia(dir_saida: str) -> Dict[str, Tuple[str, Optional[dt.date]]]:
    """Procura 'Enviado?' marcado nas planilhas do dia. Devolve {id: (valor, data da planilha)}."""
    from openpyxl import load_workbook

    resultado = {}
    for caminho in sorted(glob.glob(os.path.join(dir_saida, "vagas_*.xlsx"))):
        m = re.search(r"vagas_(\d{4}-\d{2}-\d{2})", os.path.basename(caminho))
        data = para_data(m.group(1)) if m else None
        try:
            ws = load_workbook(caminho, read_only=False).active
        except Exception:  # noqa: BLE001  planilha aberta ou corrompida: ignora
            continue
        titulos = [str(c.value or "").strip() for c in ws[1]]
        if "Enviado?" not in titulos or "Link para candidatura" not in titulos:
            continue
        i_env, i_link = titulos.index("Enviado?"), titulos.index("Link para candidatura")
        i_tit = titulos.index("Vaga") if "Vaga" in titulos else None
        i_emp = titulos.index("Empresa") if "Empresa" in titulos else None
        for linha in ws.iter_rows(min_row=2):
            enviado = normalizar_enviado(linha[i_env].value)
            if not enviado:
                continue
            celula = linha[i_link]
            url = (celula.hyperlink.target if celula.hyperlink else "") or str(celula.value or "")
            titulo = linha[i_tit].value if i_tit is not None else ""
            empresa = linha[i_emp].value if i_emp is not None else ""
            resultado[id_da_vaga(url, titulo or "", empresa or "")] = (enviado, data)
    return resultado


def ler_historico_adaptador(pasta_adaptador: str) -> List[dict]:
    caminho = os.path.join(pasta_adaptador, "historico.csv")
    if not os.path.exists(caminho):
        return []
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            with open(caminho, encoding=codificacao, newline="") as f:
                linhas = list(csv.DictReader(f, delimiter=";"))
            break
        except UnicodeDecodeError:
            continue
    else:
        return []

    adaptacoes = []
    for l in linhas:
        empresa = (l.get("Empresa") or "").strip()
        cargo = (l.get("Cargo") or "").strip()
        # Leituras que falharam (vaga nao reconhecida) nao entram no controle
        if not empresa or not cargo or normalizar(empresa).startswith("nao inform"):
            continue
        try:
            compat = int(float(str(l.get("Compat. adaptado IA") or "").replace(",", ".")))
        except ValueError:
            compat = None
        try:
            alertas = int(l.get("Alertas") or 0)
        except ValueError:
            alertas = 0
        adaptacoes.append({
            "data": para_data(l.get("Data")), "empresa": empresa, "cargo": cargo,
            "versao": (l.get("Versão") or l.get("Versao") or "").strip(),
            "compat": compat, "alertas": alertas, "link": (l.get("Link") or "").strip(),
        })
    return adaptacoes


# ---------------------------------------------------------------------------
# Montagem
# ---------------------------------------------------------------------------

def _encontrar_vaga(adaptacao: dict, linhas: Dict[str, dict]) -> Optional[str]:
    """Liga uma adaptacao a uma vaga do controle: pelo link ou por empresa + cargo."""
    if "linkedin.com" in adaptacao["link"]:
        ident = id_da_vaga(adaptacao["link"])
        if ident in linhas:
            return ident
    melhor, melhor_nota = None, 0.0
    for ident, l in linhas.items():
        if not mesma_empresa(adaptacao["empresa"], l["empresa"]):
            continue
        nota = semelhanca_titulo(adaptacao["cargo"], l["titulo"])
        if nota > melhor_nota:
            melhor, melhor_nota = ident, nota
    return melhor if melhor_nota >= 0.6 else None


def montar_linhas(historico: Dict[str, dict], palavras: Optional[PalavrasChave],
                  usuario: Dict[str, dict], existentes: Dict[str, dict],
                  enviados: Dict[str, Tuple[str, Optional[dt.date]]],
                  adaptacoes: List[dict], hoje: Optional[dt.date] = None) -> Tuple[List[dict], dict]:
    hoje = hoje or dt.date.today()
    resumo = {"enviados_importados": 0, "adaptacoes_ligadas": 0, "adaptacoes_avulsas": 0}
    linhas: Dict[str, dict] = {}

    for ident, h in historico.items():
        l = {
            "id": ident, "data": para_data(h.get("data")), "titulo": h.get("titulo", ""),
            "empresa": h.get("empresa", ""), "local": h.get("local", ""),
            "link": link_limpo(h.get("link", "")), "origem": h.get("origem") or "Agente",
            "nota": h.get("nota"), "curriculo": h.get("curriculo", ""),
            "adaptado": "", "compat": None, "alertas": None,
        }
        if l["nota"] is None and palavras is not None:
            # Vagas registradas antes da versao 1.5 nao guardavam a nota: recalcula pelo titulo
            avaliada = avaliar({"titulo": l["titulo"], "empresa": l["empresa"], "local": l["local"]}, palavras)
            l["nota"], l["curriculo"] = avaliada["nota"], avaliada["curriculo"]
        linhas[ident] = l

    # Vagas incluidas a mao entram antes, para que uma adaptacao possa ser ligada a elas
    for ident, e in existentes.items():
        if ident.startswith("manual-"):
            linhas.setdefault(ident, dict(e))

    for a in adaptacoes:
        ident = _encontrar_vaga(a, linhas)
        if ident is None:
            ident = "adapt-" + hashlib.md5(
                normalizar(f"{a['empresa']}|{a['cargo']}").encode()).hexdigest()[:10]
            if ident not in linhas:
                linhas[ident] = {
                    "id": ident, "data": a["data"], "titulo": a["cargo"], "empresa": a["empresa"],
                    "local": "", "link": a["link"] if a["link"].startswith("http") else "",
                    "origem": "Adaptador", "nota": None, "curriculo": a["versao"],
                }
                resumo["adaptacoes_avulsas"] += 1
        else:
            resumo["adaptacoes_ligadas"] += 1
        l = linhas[ident]
        l["adaptado"] = "Sim"
        if a["versao"]:
            l["curriculo"] = a["versao"]  # a versao adaptada e a que foi (ou sera) enviada
        l["compat"], l["alertas"] = a["compat"], a["alertas"]

    # Nada que ja esta no controle se perde, mesmo que tenha saido das fontes
    # (historico apagado, historico.csv do Adaptador removido...)
    for ident, e in existentes.items():
        linhas.setdefault(ident, dict(e))

    for ident, l in linhas.items():
        campos = dict(usuario.get(ident) or {})
        for nome in CAMPOS_DO_USUARIO:
            campos.setdefault(nome, l.get(nome))
        campos["Enviado?"] = normalizar_enviado(campos.get("Enviado?"))
        if not campos["Enviado?"] and ident in enviados:
            campos["Enviado?"] = enviados[ident][0]
            if not campos.get("Data do envio") and enviados[ident][1]:
                campos["Data do envio"] = enviados[ident][1]
            resumo["enviados_importados"] += 1
        envio = campos.get("Data do envio")
        if campos["Enviado?"] == "Sim" and not envio:
            campos["Data do envio"] = hoje
        elif isinstance(envio, str) and para_data(envio):
            campos["Data do envio"] = para_data(envio)
        l.update(campos)

    def chave(l):
        nota = l.get("nota")
        return (l.get("data") or dt.date.min, nota if isinstance(nota, (int, float)) else 0)

    return sorted(linhas.values(), key=chave, reverse=True), resumo


# ---------------------------------------------------------------------------
# Gravacao
# ---------------------------------------------------------------------------

def _letra(nome: str) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(COL[nome])


def _faixa(nome: str) -> str:
    letra = _letra(nome)
    return f"Candidaturas!${letra}$2:${letra}${LINHAS_FORMULA}"


def gravar_controle(linhas: List[dict], caminho: str) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    fonte = Font(name="Arial", size=10)
    negrito_branco = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    negrito_escuro = Font(name="Arial", size=10, bold=True, color="1F1F1F")
    azul_escuro = PatternFill("solid", fgColor="1F3864")
    amarelo_titulo = PatternFill("solid", fgColor="FFD966")
    amarelo_celula = PatternFill("solid", fgColor="FFF6D5")
    link_fonte = Font(name="Arial", size=10, color="0563C1", underline="single")
    fina = Side(style="thin", color="D9D9D9")
    borda = Border(bottom=fina)

    wb = Workbook()

    # ---------------- Aba Candidaturas ----------------
    ws = wb.active
    ws.title = "Candidaturas"
    ws.append(TITULOS)
    for nome, largura, do_usuario in COLUNAS:
        c = ws.cell(row=1, column=COL[nome])
        c.font = negrito_escuro if do_usuario else negrito_branco
        c.fill = amarelo_titulo if do_usuario else azul_escuro
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[_letra(nome)].width = largura
    ws.cell(row=1, column=COL["Enviado?"]).comment = Comment(
        "Colunas amarelas sao suas: o agente nunca apaga o que voce escrever nelas.", "Agente de Vagas")
    ws.row_dimensions[1].height = 30

    for l in linhas:
        valores = {
            "Data": l.get("data"), "Nota": l.get("nota"), "Currículo": l.get("curriculo"),
            "Vaga": l.get("titulo"), "Empresa": l.get("empresa"), "Local": l.get("local"),
            "Link": "Abrir vaga" if l.get("link") else "", "Origem": l.get("origem"),
            "Adaptado?": l.get("adaptado"),
            "Compat. IA": (l["compat"] / 100) if isinstance(l.get("compat"), int) else None,
            "Alertas": l.get("alertas"), "ID": l.get("id"),
            **{nome: l.get(nome) for nome in CAMPOS_DO_USUARIO},
        }
        ws.append([valores[t] for t in TITULOS])
        r = ws.max_row
        for nome, _, do_usuario in COLUNAS:
            c = ws.cell(row=r, column=COL[nome])
            c.font = fonte
            c.border = borda
            if do_usuario:
                c.fill = amarelo_celula
        if l.get("link"):
            c = ws.cell(row=r, column=COL["Link"])
            c.hyperlink = l["link"]
            c.font = link_fonte
        ws.cell(row=r, column=COL["Data"]).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=COL["Data do envio"]).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=COL["Compat. IA"]).number_format = "0%"
        for nome in ("Nota", "Currículo", "Origem", "Adaptado?", "Compat. IA", "Alertas", "Enviado?"):
            ws.cell(row=r, column=COL[nome]).alignment = Alignment(horizontal="center")

    # Listas de opcoes nas colunas que voce preenche (inclusive linhas novas)
    for nome, opcoes in (("Enviado?", OPCOES_ENVIADO), ("Retorno", OPCOES_RETORNO)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "Escolha uma das opções da lista.", "Valor fora da lista"
        dv.add(f"{_letra(nome)}2:{_letra(nome)}{LINHAS_FORMULA}")
        ws.add_data_validation(dv)

    ws.column_dimensions[_letra("ID")].hidden = True  # usado so pelo agente
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{_letra('ID')}{max(ws.max_row, 2)}"

    # ---------------- Aba Indicadores ----------------
    gravar_indicadores(wb.create_sheet("Indicadores"))

    # ---------------- Aba Como usar ----------------
    gravar_instrucoes(wb.create_sheet("Como usar"))

    # Grava num arquivo temporario e so depois troca: o controle nunca fica pela metade
    temporario = caminho + ".tmp.xlsx"
    wb.save(temporario)
    try:
        if os.path.exists(caminho):
            shutil.copy2(caminho, os.path.join(os.path.dirname(caminho), NOME_BACKUP))
        os.replace(temporario, caminho)
    finally:
        if os.path.exists(temporario):
            os.remove(temporario)


def gravar_indicadores(ws) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    titulo = Font(name="Arial", size=14, bold=True, color="1F3864")
    secao = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cabec = Font(name="Arial", size=10, bold=True)
    normal = Font(name="Arial", size=10)
    nota = Font(name="Arial", size=9, italic=True, color="595959")
    fundo_secao = PatternFill("solid", fgColor="1F3864")
    fundo_cabec = PatternFill("solid", fgColor="D9E1F2")

    for letra, largura in zip("ABCDEF", (2, 44, 14, 14, 18, 16)):
        ws.column_dimensions[letra].width = largura

    data, nota_col, curric = _faixa("Data"), _faixa("Nota"), _faixa("Currículo")
    vaga, adapt = _faixa("Vaga"), _faixa("Adaptado?")
    env, data_env, ret = _faixa("Enviado?"), _faixa("Data do envio"), _faixa("Retorno")

    def positivos(*criterios: str) -> str:
        extra = "".join("," + c for c in criterios)
        return "+".join(f'COUNTIFS({env},"Sim",{ret},"{r}"{extra})' for r in RETORNOS_POSITIVOS)

    linha = [1]

    def escrever(texto, valor=None, formato=None, fonte=normal):
        r = linha[0]
        ws.cell(row=r, column=2, value=texto).font = fonte
        if valor is not None:
            c = ws.cell(row=r, column=3, value=valor)
            c.font = fonte
            c.alignment = Alignment(horizontal="center")
            if formato:
                c.number_format = formato
        linha[0] += 1
        return r

    def cabecalho_secao(texto):
        linha[0] += 1
        r = linha[0]
        for col in range(2, 7):
            ws.cell(row=r, column=col).fill = fundo_secao
        ws.cell(row=r, column=2, value=texto).font = secao
        linha[0] += 1

    def tabela(cabecalhos, linhas_tabela, formatos):
        r = linha[0]
        for i, t in enumerate(cabecalhos):
            c = ws.cell(row=r, column=2 + i, value=t)
            c.font, c.fill = cabec, fundo_cabec
            c.alignment = Alignment(horizontal="center" if i else "left", wrap_text=True)
        for valores in linhas_tabela:
            r += 1
            for i, v in enumerate(valores):
                c = ws.cell(row=r, column=2 + i, value=v)
                c.font = normal
                if i:
                    c.alignment = Alignment(horizontal="center")
                if formatos[i]:
                    c.number_format = formatos[i]
        linha[0] = r + 1

    escrever("Indicadores da busca de emprego", fonte=titulo)
    escrever("Calculados pela aba Candidaturas. Atualizam sozinhos quando você preenche Enviado? e Retorno.",
             fonte=nota)

    cabecalho_secao("Funil geral")
    r_vagas = escrever("Vagas no controle", f"=COUNTA({vaga})")
    r_adapt = escrever("Currículos adaptados", f'=COUNTIF({adapt},"Sim")')
    r_env = escrever("Candidaturas enviadas", f'=COUNTIF({env},"Sim")')
    escrever("Taxa de envio (enviadas ÷ vagas)", f"=IFERROR(C{r_env}/C{r_vagas},0)", "0.0%")
    r_pos = escrever("Retornos positivos (entrevista, teste ou proposta)", "=" + positivos())
    escrever("Taxa de retorno (positivos ÷ enviadas)", f"=IFERROR(C{r_pos}/C{r_env},0)", "0.0%")
    r_rec = escrever("Recusados", f'=COUNTIFS({env},"Sim",{ret},"Recusado")')
    r_sem = escrever("Sem retorno", f'=COUNTIFS({env},"Sim",{ret},"Sem retorno")')
    escrever("Aguardando resposta", f"=C{r_env}-C{r_pos}-C{r_rec}-C{r_sem}")
    escrever("Adaptadas que foram enviadas", f'=COUNTIFS({env},"Sim",{adapt},"Sim")')
    _ = r_adapt

    cabecalho_secao("Qual currículo traz mais retorno?")
    linhas_tab = []
    for letra in ("A", "B"):
        r = linha[0] + 1 + len(linhas_tab)
        linhas_tab.append([
            f"Currículo {letra}", f'=COUNTIF({curric},"{letra}")',
            f'=COUNTIFS({curric},"{letra}",{env},"Sim")',
            "=" + positivos(f'{curric},"{letra}"'),
            f"=IFERROR(E{r}/D{r},0)",
        ])
    tabela(["Currículo", "Vagas", "Enviadas", "Retornos positivos", "Taxa de retorno"],
           linhas_tab, [None, None, None, None, "0.0%"])

    cabecalho_secao("Adaptar o currículo faz diferença?")
    linhas_tab = []
    for rotulo, criterio in (("Enviadas com currículo adaptado", '"Sim"'),
                             ("Enviadas sem adaptar", '"<>Sim"')):
        r = linha[0] + 1 + len(linhas_tab)
        linhas_tab.append([
            rotulo, f"=COUNTIFS({env},\"Sim\",{adapt},{criterio})",
            "=" + positivos(f"{adapt},{criterio}"), f"=IFERROR(D{r}/C{r},0)",
        ])
    tabela(["Grupo", "Enviadas", "Retornos positivos", "Taxa de retorno"],
           linhas_tab, [None, None, None, "0.0%"])

    cabecalho_secao("A nota do agente acerta? (retorno por faixa de nota)")
    linhas_tab = []
    for rotulo, minimo, maximo in (("Nota até 3", 0, 3), ("Nota 4 a 5", 4, 5),
                                   ("Nota 6 a 7", 6, 7), ("Nota 8 ou mais", 8, 999)):
        r = linha[0] + 1 + len(linhas_tab)
        faixa = f'{nota_col},">={minimo}",{nota_col},"<={maximo}"'
        linhas_tab.append([
            rotulo, f"=COUNTIFS({faixa})", f'=COUNTIFS({env},"Sim",{faixa})',
            "=" + positivos(faixa), f"=IFERROR(E{r}/D{r},0)",
        ])
    tabela(["Faixa", "Vagas", "Enviadas", "Retornos positivos", "Taxa de retorno"],
           linhas_tab, [None, None, None, None, "0.0%"])

    cabecalho_secao("Últimas 8 semanas")
    linhas_tab = []
    for k in range(8):
        r = linha[0] + 1 + len(linhas_tab)
        linhas_tab.append([
            f"=TODAY()-WEEKDAY(TODAY(),2)+1-7*{k}",
            f'=COUNTIFS({data},">="&B{r},{data},"<"&B{r}+7)',
            f'=COUNTIFS({env},"Sim",{data_env},">="&B{r},{data_env},"<"&B{r}+7)',
        ])
    tabela(["Semana que começa em", "Vagas novas", "Candidaturas enviadas"],
           linhas_tab, ["DD/MM/YYYY", None, None])
    escrever("Retorno positivo = Entrevista, Teste/Case ou Proposta na coluna Retorno.", fonte=nota)
    escrever("A semana de envio usa a coluna 'Data do envio'.", fonte=nota)
    ws.sheet_view.showGridLines = False


def gravar_instrucoes(ws) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 90
    ws.sheet_view.showGridLines = False
    titulo = Font(name="Arial", size=14, bold=True, color="1F3864")
    negrito = Font(name="Arial", size=10, bold=True)
    normal = Font(name="Arial", size=10)
    amarelo = PatternFill("solid", fgColor="FFF6D5")

    ws["B1"] = "Como usar o controle de candidaturas"
    ws["B1"].font = titulo
    itens = [
        ("O que é", "Uma lista única com todas as vagas que o agente já mostrou, as que você adaptou "
                    "e as que você enviou. O agente atualiza esta planilha todo dia, na execução das 9h."),
        ("Colunas amarelas", "São suas: Enviado?, Data do envio, Retorno e Observações. "
                             "O agente nunca apaga o que você escrever nelas."),
        ("Enviado?", "Escolha Sim quando se candidatar. Pode continuar marcando 'sim' na planilha do dia: "
                     "o agente traz a marcação para cá na próxima execução."),
        ("Data do envio", "Se ficar em branco, o agente preenche com a data em que encontrou o 'Sim'."),
        ("Retorno", "Atualize quando a empresa responder: Aguardando, Sem retorno, Entrevista, "
                    "Teste/Case, Proposta ou Recusado."),
        ("Demais colunas", "Preenchidas pelo agente. Adaptado? e Compat. IA vêm do Adaptador de Currículo."),
        ("Vaga encontrada fora do agente", "Acrescente uma linha no fim da aba Candidaturas, com Data, "
                                           "Vaga e Empresa. Deixe a coluna ID em branco."),
        ("Planilha aberta às 9h", "Se o controle estiver aberto no Excel na hora da execução, ele é atualizado "
                                  "na execução seguinte. Nada se perde."),
        ("Cópia de segurança", "Antes de cada atualização, a versão anterior é salva como "
                               "controle_candidaturas_backup.xlsx, na mesma pasta."),
    ]
    r = 3
    for chave, texto in itens:
        ws.cell(row=r, column=2, value=chave).font = negrito
        c = ws.cell(row=r, column=3, value=texto)
        c.font = normal
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="Exemplo de linha preenchida").font = negrito
    r += 1
    exemplo = [("Vaga", "Analista de Processos Pleno"), ("Empresa", "Empresa Exemplo"),
               ("Enviado?", "Sim"), ("Data do envio", "15/10/2026"), ("Retorno", "Entrevista"),
               ("Observações", "Entrevista marcada para 22/10, 14h, com a gestora da área.")]
    for chave, valor in exemplo:
        ws.cell(row=r, column=2, value=chave).font = normal
        c = ws.cell(row=r, column=3, value=valor)
        c.font = normal
        if chave in CAMPOS_DO_USUARIO:
            c.fill = amarelo
        r += 1


def atualizar_controle(dir_saida: str, historico: Dict[str, dict], pasta_adaptador: str,
                       palavras: Optional[PalavrasChave] = None,
                       hoje: Optional[dt.date] = None) -> dict:
    """Reconstroi o controle. Devolve um resumo; 'erro' preenchido se nao conseguiu gravar."""
    caminho = os.path.join(dir_saida, NOME_ARQUIVO)
    os.makedirs(dir_saida, exist_ok=True)
    try:
        usuario, existentes = ler_controle_existente(caminho)
    except Exception as e:  # noqa: BLE001
        return {"erro": f"nao consegui ler o controle atual ({e}). Ele foi mantido como estava."}
    linhas, resumo = montar_linhas(
        historico, palavras, usuario, existentes,
        ler_enviados_das_planilhas_do_dia(dir_saida),
        ler_historico_adaptador(pasta_adaptador), hoje=hoje)
    try:
        gravar_controle(linhas, caminho)
    except PermissionError:
        resumo["erro"] = ("o controle esta aberto no Excel. Ele sera atualizado na proxima execucao "
                          "(nada se perde).")
    resumo.update(caminho=caminho, linhas=len(linhas),
                  enviadas=sum(1 for l in linhas if l.get("Enviado?") == "Sim"))
    return resumo
