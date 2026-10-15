"""Execucao completa do agente, com o Gmail simulado pelo e-mail de exemplo."""
import json
import os
import shutil

from openpyxl import load_workbook

import agente_vagas
import configuracao
from conftest import PASTA_COLETOR


def preparar(tmp_path, monkeypatch, html_exemplo, extra_config=""):
    pasta_agente = tmp_path / "agente"
    pasta_agente.mkdir()
    shutil.copy(os.path.join(PASTA_COLETOR, "palavras_chave.txt"), pasta_agente)
    saida = tmp_path / "saida"
    (pasta_agente / "config.txt").write_text(
        "email = alertas@exemplo.com\nsenha = abcdefghijklmnop\nnota_minima = 4\n"
        f"pasta_saida = {saida}\npasta_adaptador = {tmp_path / 'Adaptador'}\nia_titulos = nao\n"
        + extra_config, encoding="utf-8")
    monkeypatch.setattr(configuracao, "PASTA_AGENTE", str(pasta_agente))
    monkeypatch.setattr(agente_vagas, "ler_emails",
                        lambda cfg, pasta, dias: [{"assunto": "Alerta", "html": html_exemplo}])
    return saida


def test_execucao_completa(tmp_path, monkeypatch, html_exemplo, capsys):
    saida = preparar(tmp_path, monkeypatch, html_exemplo)
    assert agente_vagas.main([]) == 0

    planilhas = list(saida.glob("vagas_*.xlsx"))
    assert len(planilhas) == 1
    ws = load_workbook(planilhas[0]).active
    titulos = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
    assert titulos == ["Analista de Processos e Melhoria Contínua", "Analista de Requisitos Sênior",
                       "Analista de Sistemas Sênior"]

    historico = json.loads((saida / "vagas_vistas.json").read_text(encoding="utf-8"))
    assert set(historico) == {"li-1000000001", "li-1000000002", "li-1000000006"}  # 7 = anuncio repetido
    assert historico["li-1000000001"]["nota"] == 9

    controle = load_workbook(saida / "controle_candidaturas.xlsx")
    assert controle["Candidaturas"].max_row == 4
    assert "Agente de Vagas v1.5" in capsys.readouterr().out


def test_segunda_execucao_nao_repete_vagas_nem_apaga_a_planilha(tmp_path, monkeypatch, html_exemplo, capsys):
    saida = preparar(tmp_path, monkeypatch, html_exemplo)
    agente_vagas.main([])
    planilha = next(saida.glob("vagas_*.xlsx"))
    antes = planilha.read_bytes()

    assert agente_vagas.main([]) == 0
    assert "Nenhuma vaga nova hoje" in capsys.readouterr().out
    assert planilha.read_bytes() == antes


def test_refazer_nao_apaga_o_historico(tmp_path, monkeypatch, html_exemplo):
    saida = preparar(tmp_path, monkeypatch, html_exemplo)
    saida.mkdir()
    (saida / "vagas_vistas.json").write_text(json.dumps({
        "li-999": {"data": "2026-09-01", "titulo": "Vaga antiga", "empresa": "X", "link": ""},
        "li-1000000001": {"data": "2026-10-01", "titulo": "Analista de Processos e Melhoria Contínua",
                          "empresa": "Empresa Alfa Ltda", "link": ""}}), encoding="utf-8")
    agente_vagas.main(["--refazer"])
    historico = json.loads((saida / "vagas_vistas.json").read_text(encoding="utf-8"))
    assert "li-999" in historico and "li-1000000002" in historico
    assert historico["li-1000000001"]["data"] == "2026-10-01"  # --refazer nao muda a data original
    assert historico["li-1000000001"]["nota"] == 9


def test_config_faltando_retorna_erro(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(configuracao, "PASTA_AGENTE", str(tmp_path))
    assert agente_vagas.main([]) == 1
    assert "config.txt" in capsys.readouterr().out


def test_segunda_execucao_com_vagas_novas_nao_sobrescreve_a_da_manha(tmp_path, monkeypatch, html_exemplo):
    saida = preparar(tmp_path, monkeypatch, html_exemplo)
    agente_vagas.main([])
    da_manha = next(saida.glob("vagas_*.xlsx"))
    antes = da_manha.read_bytes()
    # Chega um alerta novo no mesmo dia
    novo = (html_exemplo.replace("1000000001", "2000000001").replace("Empresa Alfa", "Empresa Zeta")
            .replace("1000000006", "2000000006").replace("Épsilon", "Ômega"))
    monkeypatch.setattr(agente_vagas, "ler_emails", lambda cfg, pasta, dias: [{"assunto": "A", "html": novo}])
    agente_vagas.main([])
    assert da_manha.read_bytes() == antes
    assert len(list(saida.glob("vagas_*.xlsx"))) == 2


def test_anuncio_repetido_com_outro_numero_e_ignorado(tmp_path, monkeypatch, html_exemplo):
    from agente_vagas import separar_novas
    import datetime as dt
    base = {"titulo": "Senior Functional Tester", "empresa": "EPAM Systems", "local": ""}
    brutas = [dict(base, link=f"https://www.linkedin.com/jobs/view/{n}/") for n in (1, 2, 3)]
    assert len(separar_novas(brutas, {}, dt.date(2026, 10, 14))) == 1

    recente = {"li-9": {"data": "2026-10-01", "titulo": "Senior Functional Tester", "empresa": "EPAM Systems"}}
    antigo = {"li-9": {"data": "2026-08-01", "titulo": "Senior Functional Tester", "empresa": "EPAM Systems"}}
    assert separar_novas([dict(brutas[0])], recente, dt.date(2026, 10, 14)) == []
    assert len(separar_novas([dict(brutas[0])], antigo, dt.date(2026, 10, 14))) == 1  # mais de 30 dias: vale de novo


def test_planilha_antiga_leva_o_enviado_ao_controle_antes_de_ir_para_pode_apagar(tmp_path, monkeypatch, html_exemplo):
    from planilha_dia import COL_ENVIADO, gerar_planilha_do_dia
    saida = preparar(tmp_path, monkeypatch, html_exemplo)
    saida.mkdir()
    (saida / "vagas_vistas.json").write_text(json.dumps({"li-555": {
        "data": "2026-09-01", "titulo": "Analista de Processos", "empresa": "Antiga SA",
        "link": "https://www.linkedin.com/jobs/view/555/"}}), encoding="utf-8")
    antiga = saida / "vagas_2026-09-01.xlsx"
    gerar_planilha_do_dia([{"nota": 4, "curriculo": "A", "titulo": "Analista de Processos", "empresa": "Antiga SA",
                            "local": "", "plataforma": "LinkedIn", "termos": "",
                            "link": "https://www.linkedin.com/jobs/view/555/"}], str(antiga))
    wb = load_workbook(antiga)
    wb.active.cell(row=2, column=COL_ENVIADO).value = "sim"
    wb.save(antiga)

    assert agente_vagas.main([]) == 0
    assert not antiga.exists()
    assert (saida / "PODE_APAGAR" / "vagas_2026-09-01.xlsx").exists()
    ws = load_workbook(saida / "controle_candidaturas.xlsx")["Candidaturas"]
    titulos = [c.value for c in ws[1]]
    linha = [dict(zip(titulos, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[-1] == "li-555"][0]
    assert linha["Enviado?"] == "Sim"
