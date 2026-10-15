import datetime as dt
import os

from openpyxl import load_workbook

from controle import (COL, NOME_ARQUIVO, atualizar_controle, mesma_empresa, para_data,
                      semelhanca_titulo)
from planilha_dia import COL_ENVIADO, gerar_planilha_do_dia

HOJE = dt.date(2026, 10, 15)

HISTORICO = {
    "li-1": {"data": "2026-10-10", "titulo": "Analista de Processos Pleno", "empresa": "Empresa Alfa",
             "link": "https://www.linkedin.com/comm/jobs/view/1/?otpToken=SEGREDO", "nota": 7, "curriculo": "A",
             "local": "Salvador, Bahia"},
    "li-2": {"data": "2026-10-12", "titulo": "Analista de Requisitos", "empresa": "Beta Tecnologia",
             "link": "https://www.linkedin.com/jobs/view/2/"},  # registro antigo, sem nota
}


def linhas_do_controle(pasta):
    ws = load_workbook(os.path.join(pasta, NOME_ARQUIVO))["Candidaturas"]
    titulos = [c.value for c in ws[1]]
    return {r[COL["ID"] - 1]: dict(zip(titulos, r)) for r in ws.iter_rows(min_row=2, values_only=True)}


def escrever_historico_adaptador(pasta, linhas):
    os.makedirs(pasta, exist_ok=True)
    with open(os.path.join(pasta, "historico.csv"), "w", encoding="utf-8-sig") as f:
        f.write("Data;Empresa;Cargo;Versão;Nota local A;Nota local B;Compat. base IA;"
                "Compat. adaptado IA;Alertas;Link;Currículo gerado\n")
        for l in linhas:
            f.write(l + "\n")


def test_cria_controle_com_historico_e_nota_recalculada(tmp_path, palavras):
    resumo = atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    assert "erro" not in resumo and resumo["linhas"] == 2
    linhas = linhas_do_controle(tmp_path)
    assert linhas["li-1"]["Nota"] == 7
    assert linhas["li-2"]["Nota"] >= 3 and linhas["li-2"]["Currículo"] == "B"  # recalculada pelo titulo
    assert linhas["li-1"]["Data"].date() == dt.date(2026, 10, 10)


def test_link_do_controle_nao_guarda_tokens(tmp_path, palavras):
    atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    ws = load_workbook(tmp_path / NOME_ARQUIVO)["Candidaturas"]
    alvos = [ws.cell(row=r, column=COL["Link"]).hyperlink.target for r in range(2, ws.max_row + 1)]
    assert "https://www.linkedin.com/jobs/view/1/" in alvos
    assert not any("SEGREDO" in a for a in alvos)


def test_preserva_o_que_o_usuario_escreveu(tmp_path, palavras):
    atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    caminho = tmp_path / NOME_ARQUIVO
    wb = load_workbook(caminho)
    ws = wb["Candidaturas"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=COL["ID"]).value == "li-1":
            ws.cell(row=r, column=COL["Enviado?"]).value = "sim"
            ws.cell(row=r, column=COL["Retorno"]).value = "Entrevista"
            ws.cell(row=r, column=COL["Observações"]).value = "Falar com a gestora"
    # Vaga incluida a mao, sem ID
    ws.append({COL["Vaga"]: "Coordenador da Qualidade", COL["Empresa"]: "Gama SA", COL["Data"]: "14/10/2026"})
    wb.save(caminho)

    atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    linhas = linhas_do_controle(tmp_path)
    assert linhas["li-1"]["Enviado?"] == "Sim"
    assert linhas["li-1"]["Data do envio"].date() == HOJE  # preenchida pelo agente
    assert linhas["li-1"]["Retorno"] == "Entrevista"
    assert linhas["li-1"]["Observações"] == "Falar com a gestora"
    manual = [l for i, l in linhas.items() if str(i).startswith("manual-")]
    assert len(manual) == 1 and manual[0]["Vaga"] == "Coordenador da Qualidade"
    assert manual[0]["Origem"] == "Manual"
    assert os.path.exists(tmp_path / "controle_candidaturas_backup.xlsx")


def test_traz_o_enviado_marcado_na_planilha_do_dia(tmp_path, palavras):
    vagas = [{"nota": 7, "curriculo": "A", "titulo": "Analista de Processos Pleno", "empresa": "Empresa Alfa",
              "local": "", "plataforma": "LinkedIn", "termos": "",
              "link": "https://www.linkedin.com/comm/jobs/view/1/?trk=x"}]
    caminho = tmp_path / "vagas_2026-10-10.xlsx"
    gerar_planilha_do_dia(vagas, str(caminho))
    wb = load_workbook(caminho)
    wb.active.cell(row=2, column=COL_ENVIADO).value = "SIM"
    wb.save(caminho)

    resumo = atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    assert resumo["enviados_importados"] == 1
    linha = linhas_do_controle(tmp_path)["li-1"]
    assert linha["Enviado?"] == "Sim"
    assert linha["Data do envio"].date() == dt.date(2026, 10, 10)  # data da planilha


def test_liga_adaptacao_a_vaga_e_registra_avulsa(tmp_path, palavras):
    adaptador = tmp_path / "Adaptador"
    escrever_historico_adaptador(str(adaptador), [
        "13/10/2026 10:00;Empresa Alfa Ltda;Analista de Processos Pleno;A;20;10;60;88;0;;cv1.docx",
        "13/10/2026 11:00;Zeta Consultoria;Analista de Negócios;B;10;30;50;80;2;;cv2.docx",
        "13/10/2026 12:00;Nao informado;Nao informado;A;0;0;0;0;1;;cv3.docx",  # leitura que falhou
    ])
    resumo = atualizar_controle(str(tmp_path), HISTORICO, str(adaptador), palavras, hoje=HOJE)
    assert resumo["adaptacoes_ligadas"] == 1 and resumo["adaptacoes_avulsas"] == 1
    linhas = linhas_do_controle(tmp_path)
    assert linhas["li-1"]["Adaptado?"] == "Sim"
    assert linhas["li-1"]["Compat. IA"] == 0.88
    avulsa = [l for i, l in linhas.items() if str(i).startswith("adapt-")][0]
    assert avulsa["Empresa"] == "Zeta Consultoria" and avulsa["Currículo"] == "B"
    assert avulsa["Origem"] == "Adaptador" and avulsa["Alertas"] == 2
    assert len(linhas) == 3


def test_indicadores_tem_formulas(tmp_path, palavras):
    atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    ws = load_workbook(tmp_path / NOME_ARQUIVO)["Indicadores"]
    formulas = [c.value for linha in ws.iter_rows() for c in linha
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 30
    assert any("COUNTIFS" in f and "Entrevista" in f for f in formulas)


def test_controle_aberto_no_excel_nao_perde_nada(tmp_path, palavras, monkeypatch):
    atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    antes = (tmp_path / NOME_ARQUIVO).read_bytes()

    def bloqueado(*a, **k):
        raise PermissionError("arquivo em uso")

    monkeypatch.setattr(os, "replace", bloqueado)
    resumo = atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    assert "aberto no Excel" in resumo["erro"]
    assert (tmp_path / NOME_ARQUIVO).read_bytes() == antes
    assert not list(tmp_path.glob("*.tmp.xlsx"))


def test_auxiliares():
    assert para_data("06/10/2026 00:21") == dt.date(2026, 10, 6)
    assert para_data("2026-10-06") == dt.date(2026, 10, 6)
    assert para_data("amanha") is None
    assert mesma_empresa("Tata Consultancy Services", "TATA CONSULTANCY")
    assert not mesma_empresa("Alfa", "")
    assert semelhanca_titulo("Analista de Negócios Pleno", "Analista de negócios") == 1.0


def test_vaga_que_saiu_do_historico_continua_no_controle(tmp_path, palavras):
    atualizar_controle(str(tmp_path), HISTORICO, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    caminho = tmp_path / NOME_ARQUIVO
    wb = load_workbook(caminho)
    ws = wb["Candidaturas"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=COL["ID"]).value == "li-1":
            ws.cell(row=r, column=COL["Enviado?"]).value = "Sim"
            ws.cell(row=r, column=COL["Retorno"]).value = "Proposta"
    wb.save(caminho)

    # Historico apagado: a vaga li-1 nao vem mais das fontes
    atualizar_controle(str(tmp_path), {}, str(tmp_path / "Adaptador"), palavras, hoje=HOJE)
    linhas = linhas_do_controle(tmp_path)
    assert linhas["li-1"]["Retorno"] == "Proposta"
    assert linhas["li-1"]["Vaga"] == "Analista de Processos Pleno"
    assert linhas["li-1"]["Nota"] == 7 and linhas["li-1"]["Origem"] == "Agente"
