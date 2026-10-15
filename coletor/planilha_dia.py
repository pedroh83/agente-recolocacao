# -*- coding: utf-8 -*-
"""Planilha do dia (vagas_AAAA-MM-DD.xlsx): as vagas novas, ordenadas por nota."""
from __future__ import annotations

from typing import List

CABECALHO = ["Nota", "Currículo", "Vaga", "Empresa", "Local",
             "Plataforma", "Link para candidatura", "Termos que bateram", "Enviado?"]
LARGURAS = [7, 11, 46, 30, 24, 12, 22, 40, 11]
COL_LINK = CABECALHO.index("Link para candidatura") + 1
COL_ENVIADO = CABECALHO.index("Enviado?") + 1


def gerar_planilha_do_dia(vagas: List[dict], caminho: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Vagas do dia"
    ws.append(CABECALHO)

    azul = PatternFill("solid", fgColor="1F3864")
    for col in range(1, len(CABECALHO) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = Alignment(horizontal="center", vertical="center")

    verde = PatternFill("solid", fgColor="E2EFDA")
    amarelo = PatternFill("solid", fgColor="FFF2CC")

    for v in vagas:
        ws.append([v["nota"], v["curriculo"], v["titulo"], v["empresa"], v["local"],
                   v["plataforma"], "Candidatar-se", v["termos"], ""])
        r = ws.max_row
        link = ws.cell(row=r, column=COL_LINK)
        link.hyperlink = v["link"]
        link.font = Font(color="0563C1", underline="single")

        cor = verde if v["nota"] >= 6 else amarelo if v["nota"] >= 3 else None
        if cor:
            for col in range(1, len(CABECALHO) + 1):
                ws.cell(row=r, column=col).fill = cor

    for i, largura in enumerate(LARGURAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO))}{ws.max_row}"
    wb.save(caminho)
